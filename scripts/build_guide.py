from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path(__file__).resolve().parents[1] / "docs" / "Docker_Profile_App_Guide.docx"
NAVY = RGBColor(31, 77, 120)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(90, 99, 110)
LIGHT = "E8EEF5"
PALE = "F4F6F9"
INK = RGBColor(25, 34, 45)

def font(run, size=11, bold=False, color=INK, name="Calibri", italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + tag))
        if node is None:
            node = OxmlElement("w:" + tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table_pr = table._tbl.tblPr
    table_w = table_pr.find(qn("w:tblW"))
    if table_w is None:
        table_w = OxmlElement("w:tblW")
        table_pr.append(table_w)
    table_w.set(qn("w:w"), str(sum(widths)))
    table_w.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_width(cell, widths[i]); margins(cell); shade(cell, LIGHT)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        font(p.add_run(text), 9.5, True, NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            set_cell_width(cells[i], widths[i]); margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
            font(p.add_run(str(text)), 9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

def para(doc, text="", bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_lead and text.startswith(bold_lead):
        font(p.add_run(bold_lead), 11, True)
        font(p.add_run(text[len(bold_lead):]), 11)
    else:
        font(p.add_run(text), 11)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(.375)
    p.paragraph_format.first_line_indent = Inches(-.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), 11)

def step(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(.375)
    p.paragraph_format.first_line_indent = Inches(-.188)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    font(p.add_run(text), 11)

def code(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0); set_cell_width(cell, 9360); margins(cell, 100, 150, 100, 150); shade(cell, "F2F4F7")
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
    for n, line in enumerate(text.splitlines()):
        if n: p.add_run().add_break()
        font(p.add_run(line), 9, color=RGBColor(35, 43, 52), name="Courier New")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    font(p.add_run(text), 16 if level == 1 else 13, True, BLUE if level < 3 else NAVY)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(1)
section.left_margin = section.right_margin = Inches(1)
section.header_distance = section.footer_distance = Inches(.492)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(11); normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.25
for name, size, color, before, after in (("Heading 1",16,BLUE,18,10),("Heading 2",13,BLUE,14,7),("Heading 3",12,NAVY,10,5)):
    style = doc.styles[name]; style.font.name="Calibri"; style.font.size=Pt(size); style.font.bold=True; style.font.color.rgb=color
    style.paragraph_format.space_before=Pt(before); style.paragraph_format.space_after=Pt(after); style.paragraph_format.keep_with_next=True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
font(header.add_run("PROFILE APP  |  DOCKER LEARNING GUIDE"), 8.5, True, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
font(footer.add_run("MongoDB • Mongo Express • Node.js • Docker Compose"), 8.5, color=MUTED)

p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(36); p.paragraph_format.space_after=Pt(8)
font(p.add_run("DOCKER PROFILE APP"), 27, True, NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_after=Pt(18)
font(p.add_run("Build, run, test, persist, and publish a three-service application"), 15, color=MUTED)
add_table(doc, ["Component", "Role", "Local address"], [
    ("Profile app", "Frontend and Express API", "http://localhost:3000"),
    ("MongoDB", "Persistent document database", "127.0.0.1:27017"),
    ("Mongo Express", "Browser database administration UI", "http://localhost:8081"),
], [1800, 4560, 3000])
para(doc, "Scope: the working setup as of September 2026. Secrets are intentionally omitted. Keep the real .env file private.")

heading(doc, "1. Architecture")
code(doc, "Browser ──► localhost:3000 ──► app container ──► mongodb:27017\nBrowser ──► localhost:8081 ──► mongo-express ──► mongodb:27017\n                                      │\n                                      └── mongo-data volume (/data/db)")
para(doc, "Compose creates one private network. Containers resolve each other by service name, so both the app and Mongo Express use mongodb as the database hostname. A program running directly on the Mac uses 127.0.0.1 and MongoDB's published port instead.")

heading(doc, "2. Project files")
add_table(doc, ["File", "Purpose"], [
    ("compose.yaml", "Portable definition of the app, MongoDB, Mongo Express, network, health checks, and volumes."),
    ("compose.override.yaml", "Ignored local override that preserves this Mac's original database volumes and pinned MongoDB image."),
    ("Dockerfile", "Builds the Node.js application image."),
    ("server.js", "Serves the frontend and implements the profile API."),
    ("public/", "Browser HTML, CSS, and JavaScript."),
    (".env", "Private credentials and runtime choices. Never commit or publish it."),
    (".env.example", "Safe configuration template with placeholders."),
    ("scripts/setup-env.cjs", "Creates a new .env with matching random credentials."),
], [2400, 6960])

heading(doc, "3. Environment variables")
add_table(doc, ["Variable", "Used by", "Meaning"], [
    ("MONGO_ROOT_USERNAME", "MongoDB", "Administrator created on first initialization."),
    ("MONGO_ROOT_PASSWORD", "MongoDB", "Administrator password; must match both connection URLs."),
    ("MONGODB_URI", "Local Node.js", "Uses 127.0.0.1 when Node runs on the Mac."),
    ("MONGO_EXPRESS_URI", "App + Mongo Express containers", "Uses mongodb, the Compose service hostname."),
    ("MONGODB_DB", "App", "Database selected by the application."),
    ("MONGODB_COLLECTION", "App", "Collection selected by the application."),
    ("MONGO_EXPRESS_WEB_*", "Mongo Express", "Separate browser login for port 8081."),
    ("APP_IMAGE", "Compose", "Image tag to build locally or pull from a registry."),
], [2700, 2400, 4260])
code(doc, "# Example only — replace placeholders\nMONGODB_URI='mongodb://USER:PASSWORD@127.0.0.1:27017/?authSource=admin'\nMONGO_EXPRESS_URI='mongodb://USER:PASSWORD@mongodb:27017/?authSource=admin'")

heading(doc, "4. Start and inspect the stack")
step(doc, "Start Docker Desktop.")
step(doc, "Open a terminal in the project directory.")
step(doc, "Build the app and start all services.")
code(doc, "docker compose up -d --build\ndocker compose ps")
step(doc, "Open the profile website on port 3000 and Mongo Express on port 8081.")
para(doc, "Compose waits for MongoDB to become healthy before starting dependent services. The app health check calls GET /api/ready, which confirms database read access. GET /api/health checks process liveness separately.")

heading(doc, "5. Manual test checklist")
bullet(doc, "Open http://localhost:3000 and confirm the profile form loads.")
bullet(doc, "Enter a name, valid email, and interests; click Save.")
bullet(doc, "Refresh the page and confirm the saved values return.")
bullet(doc, "Open http://localhost:8081 and sign in with the Mongo Express web credentials from .env.")
bullet(doc, "Select the configured database and collection; confirm the demo-profile document matches the form.")
bullet(doc, "Change a value at port 3000, save it, and refresh Mongo Express to see the update.")
code(doc, "curl http://localhost:3000/api/profile\ndocker compose logs --tail=50 app mongodb mongo-express")

heading(doc, "6. Volumes and persistence")
para(doc, "A container is replaceable. A Docker volume stores data outside the container's writable layer. MongoDB mounts mongo-data at /data/db, so profiles remain after containers are removed and recreated.")
add_table(doc, ["Storage type", "Typical use", "Behavior"], [
    ("Named volume", "Databases and uploads", "Docker-managed, recognizable name, reusable across containers."),
    ("Anonymous volume", "Temporary or automatically generated persistence", "Docker generates a long name; harder to identify and reuse."),
    ("External volume", "Pre-existing data managed outside this Compose project", "Compose uses it but does not create or remove it."),
    ("Bind mount", "Live source files and configuration", "Maps a host path directly into a container."),
    ("tmpfs", "Sensitive or disposable temporary data", "Lives in memory and disappears when stopped."),
], [1900, 3400, 4060])
code(doc, "docker compose down                 # remove containers; keep data\ndocker compose up -d                  # recreate; profile should remain\ndocker compose down --volumes         # deletes managed volumes and data")
para(doc, "A volume provides persistence, not backup. Important data still needs an independent backup strategy.")

heading(doc, "7. Build locally or pull from Docker Hub")
para(doc, "The app service has both build and image. The image field names the image; build tells Compose how to create it from source. APP_IMAGE selects the name and version.")
add_table(doc, ["Goal", "Commands"], [
    ("Develop from local source", "docker compose up -d --build"),
    ("Run published version", "docker compose pull app\ndocker compose up -d --no-build"),
    ("Rebuild only the app", "docker compose up -d --build app"),
], [3000, 6360])
para(doc, "Build current source for the recovery endpoints. The historical singhania8192/profile-app:1.0 image has not been verified with the new readiness health check. A registry stores images; it does not run the website. A computer or hosting platform must still run the containers.")

heading(doc, "8. Publishing a release")
code(doc, "docker login\ndocker build -t singhania8192/profile-app:1.1 .\ndocker push singhania8192/profile-app:1.1")
para(doc, "Use versioned tags so users can choose a stable release. Publishing a multi-platform image supports both Intel/AMD and Apple Silicon machines.")
code(doc, "docker buildx create --name profile-multiplatform --driver docker-container --use\ndocker buildx build --platform linux/amd64,linux/arm64 \\\n  -t singhania8192/profile-app:1.1 --push .")

heading(doc, "9. Common problems")
add_table(doc, ["Symptom", "Cause and fix"], [
    ("document is not defined", "public/app.js is browser code. Run server.js or Compose; do not execute public/app.js with Node."),
    ("MongoServerSelectionError", "Use 127.0.0.1 for Node on the Mac and mongodb for containers. Confirm MongoDB is healthy."),
    ("Port 3000 already in use", "Stop the manually running npm process before starting the app container."),
    ("Port 8081 unavailable", "Start mongo-express and inspect its logs. Confirm its database URL uses mongodb."),
    ("Required variable missing", "Run node scripts/setup-env.cjs or add the missing variable to .env."),
    ("New password has no effect", "Initialization variables apply only to a fresh database volume; existing users live in stored data."),
], [2800, 6560])

heading(doc, "10. Safe sharing checklist")
bullet(doc, "Keep .env and compose.override.yaml ignored by Git and Docker builds.")
bullet(doc, "Share .env.example and the setup script instead of real credentials.")
bullet(doc, "Publish the Docker image publicly for anonymous pulls.")
bullet(doc, "Publish source separately so others can inspect, modify, and build it.")
bullet(doc, "Choose a source-code license before inviting broad reuse; repository visibility alone is not a license.")

heading(doc, "11. Failure and recovery lab")
para(doc, "The separate lab uses port 3001 and its own volumes. It does not load the normal Compose override or use the existing database. Follow docs/recovery-lab.md for the full written walkthrough; no video is required.")
code(doc, "docker compose -p profile-recovery-lab -f compose.lab.yaml up -d --build --wait\nnpm run lab:test")
para(doc, "The test stops MongoDB, verifies Node and the webpage stay available, checks readiness and saves fail with HTTP 503, then starts MongoDB and measures recovery. It also replaces the database container and verifies the saved profile survives without restarting the app. See docs/recovery-results.md for observed results and docs/linkedin-post.md for the post draft.")
para(doc, "Health checks detect failure; the operator restores MongoDB in this exercise. Previously stored data remains in a volume, but failed saves require a retry. Volumes provide persistence, not backups.")

heading(doc, "12. Recommended next exercise")
para(doc, "Extend the fixed demo profile into a complete CRUD application: create multiple profiles, list them, edit a selected profile, and delete with confirmation. Then release the change as image version 1.1. This connects Node.js and MongoDB development with the Docker release workflow.")

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
